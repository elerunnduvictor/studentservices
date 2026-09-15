"""
Build supabase/tech-bugs.sql from the TS Product Tracker workbook.

    python tools/build-tech-bugs.py                      # "TS Product Tracker.xlsx"
    python tools/build-tech-bugs.py "path/to/file.xlsx"

Then run supabase/tech-bugs.sql in the Supabase SQL editor. That is the whole
weekly update: the Emerging Issues page reads the table, so nothing has to be
deployed when the list changes.

── What it reads ─────────────────────────────────────────────────────────────

Every VISIBLE sheet whose row 2 is a header starting "Priority" is one
product's bug tracker — Admissions-EE, Finance, Canvas, … — and becomes one
compartment on the page. Hidden sheets are the tracker's own superseded copies
("OLD …", "…_og", "….expired", "P-R Old") and a KB list; they are skipped and
named in the output so nothing is dropped silently. "Product Owners" is not a
bug list and is skipped too.

A tracker sheet has up to four sections, each introduced by a title row and
its own header row:

    (rows 3…)                   active     the backlog — what the page opens on
    Resolved Bugs This Week     resolved
    Closed Bugs                 closed
    Removed Bugs                removed

Columns are found by their header text, never by position, and per section:
in the last three "Latest Status" is headed "Reason for Resolution / Closure /
Removal" and "Weighted Score" is headed "Priority Estimator Score". Only the
ten columns the page shows are copied — Priority, Discovered, Bug / ADO, Scope,
Summary, Owner, Updated ETA, Workaround, Latest Status, Weighted Score.

Finance ends with "Important Notes for Finance Procedures" — two procedures,
not bugs. They are copied to tech_bug_notes and shown in that compartment.

── The history ───────────────────────────────────────────────────────────────

tech_bugs holds this week's list and nothing else: each run replaces it. The
page's line graphs need the weeks before too, so each run also writes one row
per product per section into tech_bug_history — how many bugs, how many of
them scored, and their total weighted score — dated by the day the workbook was
saved. Earlier weeks are never touched. Re-running the same week replaces that
week's rows rather than adding a second set, so a file run twice cannot draw a
point twice.

The history begins with the first workbook run through this. It is not
reconstructed backwards: the tracker keeps no record of what its lists looked
like on past dates, and a line drawn from guesses would be worse than a line
that starts late.

── What it refuses ───────────────────────────────────────────────────────────

It stops with an error, and writes nothing, if a tracker sheet is missing one
of the ten columns, if an active bug has no priority or no Bug / ADO, or if a
ticket number in a Bug / ADO cell would not survive being split into ticket
chips and a title. A wrong page is worse than an old one.
"""

import datetime
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

REPO = Path(__file__).resolve().parent.parent
DEFAULT_BOOK = REPO / "TS Product Tracker.xlsx"
OUT = REPO / "supabase" / "tech-bugs.sql"

SECTION_TITLES = {
    "resolved bugs this week": "resolved",
    "closed bugs": "closed",
    "removed bugs": "removed",
}

# Header text (normalised) -> the field it fills.
HEADERS = {
    "priority": "priority",
    "discovered": "discovered",
    "bug / ado": "bug",
    "scope": "scope",
    "summary": "summary",
    "owner": "owner",
    "updated eta": "updated_eta",
    "workaround": "workaround",
    "latest status": "latest_status",
    "reason for resolution": "latest_status",
    "reason for closure": "latest_status",
    "reason for removal": "latest_status",
    "weighted score": "weighted_score",
    "priority estimator score": "weighted_score",
}
FIELDS = ["priority", "discovered", "bug", "scope", "summary", "owner",
          "updated_eta", "workaround", "latest_status", "weighted_score"]

# How each tracker is named on the page. The sheets' own title rows cannot be
# used: Canvas's and Transcripts' both read "Planning/Registration Bug
# Tracker", copied from the sheet they were made from. A new sheet that is not
# listed here still appears, under its tab name, and the build says so.
LABELS = {
    "Admissions-EE": "Admissions / EE",
    "Finance": "Finance",
    "Canvas": "Canvas",
    "Portal-Okta": "Student Portal / Okta",
    "Companion": "Companion",
    "EC3": "EnglishConnect 3",
    "Planning-Registration": "Planning / Registration",
    "CRM Workspace": "CRM Workspace",
    "Gatherings": "Gatherings",
    "Transcripts": "Transcripts",
}

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


class BuildError(Exception):
    pass


def norm(v):
    return " ".join(str(v).replace("\xa0", " ").split()).lower() if v is not None else ""


def text(v):
    """A cell as the page should show it. Dates become ISO days; whole numbers
    lose the ".0"; text keeps its line breaks but not the Windows \r or the
    non-breaking spaces Excel likes to leave behind."""
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (int, float)):
        return str(v)
    s = str(v).replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    s = "\n".join(line.rstrip() for line in s.split("\n")).strip()
    return s or None


def as_date(v, captured):
    """A real day from a cell, or None. Written dates ("Aug 6th", "Sept 10th")
    carry no year; they take the tracker's, or the year before if that would
    put them after the day the tracker was saved. Only used for ordering and
    for "N days ago" — the page always shows the cell as it was written."""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if not isinstance(v, str):
        return None
    s = v.strip()
    m = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", s)
    if m:
        try:
            return datetime.date(int(m.group(3)), int(m.group(1)), int(m.group(2)))
        except ValueError:
            return None
    m = re.match(r"^([A-Za-z]{3,9})\.?\s+(\d{1,2})(?:st|nd|rd|th)?(?:,?\s*(\d{4}))?$", s)
    if m and m.group(1)[:3].lower() in MONTHS:
        month, day = MONTHS[m.group(1)[:3].lower()], int(m.group(2))
        year = int(m.group(3)) if m.group(3) else captured.year
        try:
            d = datetime.date(year, month, day)
        except ValueError:
            return None
        if not m.group(3) and d > captured:
            d = d.replace(year=year - 1)
        return d
    return None


def eta_date(v):
    """An ETA only counts as a date if it is one, or if the text holds a full
    m/d/yyyy — "Digital Ops: 9/18/2026". "Assigned back to ICS on 9/11" is when
    something happened, not when it will be done, and has no year: not a date."""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, str):
        m = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", v)
        if m:
            try:
                return datetime.date(int(m.group(3)), int(m.group(1)), int(m.group(2)))
            except ValueError:
                return None
    return None


def number(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and re.fullmatch(r"\s*-?\d+(\.\d+)?\s*", v):
        return float(v)
    return None


# One or more ticket numbers at the start of the cell, with the vendor that owns
# them in brackets. Without a vendor, a number only counts as a ticket if a line
# break, " - " or the end of the cell follows — so a title that happens to begin
# with a year is left alone.
REF = re.compile(
    r"^[ \t]*(\d{4,7}(?:[ \t]*,[ \t]*\d{4,7})*)[ \t]*"
    r"(?:\(([^)\n]{1,40})\)|(?=\n|$|-[ \t]))")


def parse_bug(raw):
    """Split "930545 (ICS)\\n20524 (Ellucian)\\nSame Section Registered Twice"
    into ticket chips and a title. Proved lossless below: every ticket-length
    number in the cell ends up in a chip or in the title, or the build stops."""
    refs, rest = [], raw
    while True:
        m = REF.match(rest)
        if not m:
            break
        vendor = (m.group(2) or "").strip() or None
        for i in m.group(1).split(","):
            refs.append({"id": i.strip(), "vendor": vendor})
        rest = rest[m.end():]
        rest = re.sub(r"^\s*(?:-[ \t]+)?", "", rest)
    title = " ".join(rest.split()) or None

    found = set(re.findall(r"\d{4,7}", raw))
    kept = {r["id"] for r in refs} | set(re.findall(r"\d{4,7}", title or ""))
    if found != kept:
        raise BuildError(f"Bug / ADO {raw!r}: tickets {sorted(found - kept)} would be lost")
    if not refs and not title:
        raise BuildError(f"Bug / ADO {raw!r}: nothing to show")
    return refs, title


def header_map(ws, row):
    cols, unknown = {}, []
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(row, c).value)
        if not h:
            continue
        if h in HEADERS:
            cols.setdefault(HEADERS[h], c)
        elif c <= 15:
            unknown.append(h)
    return cols, unknown


def section_title(ws, r):
    for c in range(1, 4):
        h = norm(ws.cell(r, c).value)
        if h in SECTION_TITLES:
            return SECTION_TITLES[h]
    return None


def notes_start(ws, r):
    for c in range(1, 16):
        v = ws.cell(r, c).value
        if isinstance(v, str) and norm(v).startswith("important notes"):
            return c, text(v)
    return None


def read_sheet(ws, order, captured):
    active_cols, unknown = header_map(ws, 2)
    missing = [f for f in FIELDS if f not in active_cols]
    if missing:
        raise BuildError(f"{ws.title}: header row 2 has no column for {missing}")

    bugs, notes = [], []
    section, cols = "active", active_cols
    r = 3
    while r <= ws.max_row:
        sec = section_title(ws, r)
        if sec:
            section = sec
            cols, _ = header_map(ws, r + 1)
            # A section whose header left a column blank uses the active
            # layout for it; the four sections share one set of columns.
            for f in FIELDS:
                if f not in cols and not norm(ws.cell(r + 1, active_cols[f]).value):
                    cols[f] = active_cols[f]
            gone = [f for f in FIELDS if f not in cols]
            if gone:
                raise BuildError(f"{ws.title} / {sec}: header has no column for {gone}")
            r += 2
            continue

        ns = notes_start(ws, r)
        if ns:
            heading = ns[1]
            titles = {c: text(ws.cell(r + 1, c).value) for c in range(1, ws.max_column + 1)
                      if text(ws.cell(r + 1, c).value)}
            for i, (c, title) in enumerate(sorted(titles.items())):
                body = "\n\n".join(t for t in (text(ws.cell(rr, c).value)
                                               for rr in range(r + 2, ws.max_row + 1)) if t)
                notes.append({"heading": heading, "title": title, "body": body or None,
                              "note_order": i + 1})
            break

        cells = {f: ws.cell(r, c).value for f, c in cols.items()}
        if all(text(v) is None for v in cells.values()) or norm(cells["bug"]) == "bug / ado":
            r += 1
            continue

        raw_bug = text(cells["bug"])
        if section == "active":
            if text(cells["priority"]) is None:
                raise BuildError(f"{ws.title} row {r}: an active bug with no Priority")
            if raw_bug is None:
                raise BuildError(f"{ws.title} row {r}: an active bug with no Bug / ADO")
        refs, title = parse_bug(raw_bug) if raw_bug else ([], None)

        bugs.append({
            "product": ws.title,
            "product_label": LABELS.get(ws.title, ws.title.replace("-", " / ")),
            "product_order": order,
            "section": section,
            "row_order": r,
            "priority": text(cells["priority"]),
            "discovered": text(cells["discovered"]),
            "discovered_on": as_date(cells["discovered"], captured),
            "bug": raw_bug,
            "bug_title": title,
            "bug_refs": refs,
            "scope": text(cells["scope"]),
            "summary": text(cells["summary"]),
            "owner": text(cells["owner"]),
            "updated_eta": text(cells["updated_eta"]),
            "eta_on": eta_date(cells["updated_eta"]),
            "workaround": text(cells["workaround"]),
            "latest_status": text(cells["latest_status"]),
            "weighted_score": text(cells["weighted_score"]),
            "score": number(cells["weighted_score"]),
        })
        r += 1
    return bugs, notes


def sql(v):
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(int(v)) if float(v).is_integer() else repr(v)
    if isinstance(v, datetime.date):
        return f"date '{v.isoformat()}'"
    if isinstance(v, (list, dict)):
        return "'" + json.dumps(v, ensure_ascii=False).replace("'", "''") + "'::jsonb"
    return "'" + str(v).replace("'", "''") + "'"


SCHEMA = """\
-- ── the tables ─────────────────────────────────────────────────────────────
create table if not exists public.tech_bugs (
  product         text    not null,   -- the workbook tab, e.g. 'Admissions-EE'
  product_label   text    not null,   -- how the page names it
  product_order   integer not null,   -- the tab's position in the workbook
  section         text    not null check (section in ('active', 'resolved', 'closed', 'removed')),
  row_order       integer not null,   -- the row on its sheet, which is its order
  priority        text,
  discovered      text,               -- as written; ISO day when the cell was a date
  discovered_on   date,               -- the same, as a day, where one can be read
  bug             text,               -- the Bug / ADO cell as written
  bug_title       text,
  bug_refs        jsonb   not null default '[]',   -- [{id, vendor}]
  scope           text,
  summary         text,
  owner           text,
  updated_eta     text,               -- as written
  eta_on          date,               -- only when the ETA is an actual day
  workaround      text,
  latest_status   text,               -- or the reason it was resolved/closed/removed
  weighted_score  text,               -- as written: '140', 'N/A', or empty
  score           numeric,            -- the same, where it is a number
  captured_on     date    not null,   -- the day the workbook was saved
  primary key (product, section, row_order)
);

create table if not exists public.tech_bug_notes (
  product        text    not null,
  product_label  text    not null,
  heading        text,
  title          text    not null,
  body           text,
  note_order     integer not null,
  captured_on    date    not null,
  primary key (product, note_order)
);

-- ── who may read them ──────────────────────────────────────────────────────
-- Whoever may read the Emerging Issues register: the same function its
-- policy uses (emerging-issues-partners.sql), so the two tabs of one page can
-- never disagree about who is let in. Read-only for everyone — the only way
-- rows change is this file.
alter table public.tech_bugs      enable row level security;
alter table public.tech_bug_notes enable row level security;

drop policy if exists "tech_bugs_select" on public.tech_bugs;
create policy "tech_bugs_select" on public.tech_bugs
  for select to authenticated using (public.hub_sees_emerging_issues());

drop policy if exists "tech_bug_notes_select" on public.tech_bug_notes;
create policy "tech_bug_notes_select" on public.tech_bug_notes
  for select to authenticated using (public.hub_sees_emerging_issues());

revoke all on public.tech_bugs, public.tech_bug_notes from public, anon, authenticated;
grant select on public.tech_bugs, public.tech_bug_notes to authenticated;

-- ── the weeks behind it ────────────────────────────────────────────────────
-- One row per workbook, per product, per section. What the page's line graphs
-- are drawn from. Only this week's rows are replaced by a run; every earlier
-- week stays exactly as it was recorded.
create table if not exists public.tech_bug_history (
  captured_on    date    not null,   -- the day that week's workbook was saved
  product        text    not null,
  product_label  text    not null,
  product_order  integer not null,
  section        text    not null check (section in ('active', 'resolved', 'closed', 'removed')),
  bugs           integer not null,   -- how many bugs the section listed
  scored         integer not null,   -- how many of them carried a numeric score
  score_total    numeric not null,   -- their scores added together
  top_score      numeric,            -- the highest of them
  primary key (captured_on, product, section)
);

alter table public.tech_bug_history enable row level security;
drop policy if exists "tech_bug_history_select" on public.tech_bug_history;
create policy "tech_bug_history_select" on public.tech_bug_history
  for select to authenticated using (public.hub_sees_emerging_issues());
revoke all on public.tech_bug_history from public, anon, authenticated;
grant select on public.tech_bug_history to authenticated;
"""

BUG_COLS = ["product", "product_label", "product_order", "section", "row_order", "priority",
            "discovered", "discovered_on", "bug", "bug_title", "bug_refs", "scope", "summary",
            "owner", "updated_eta", "eta_on", "workaround", "latest_status", "weighted_score",
            "score", "captured_on"]
NOTE_COLS = ["product", "product_label", "heading", "title", "body", "note_order", "captured_on"]


def main():
    book = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_BOOK
    if not book.exists():
        sys.exit(f"No workbook at {book}")
    wb = openpyxl.load_workbook(book, data_only=True)

    saved = wb.properties.modified
    if saved:
        # Stored in UTC; the day it was saved is the day where it was saved.
        captured = saved.replace(tzinfo=datetime.timezone.utc).astimezone().date()
    else:
        captured = datetime.date.fromtimestamp(book.stat().st_mtime)

    trackers, skipped = [], []
    for ws in wb.worksheets:
        is_tracker = norm(ws["A2"].value) == "priority"
        if ws.sheet_state != "visible":
            skipped.append(f"{ws.title} (hidden)")
        elif not is_tracker:
            skipped.append(f"{ws.title} (not a bug list)")
        else:
            trackers.append(ws)
    if not trackers:
        raise BuildError("no tracker sheets found")

    bugs, notes = [], []
    for i, ws in enumerate(trackers, 1):
        b, n = read_sheet(ws, i, captured)
        bugs += b
        for note in n:
            note.update(product=ws.title,
                        product_label=LABELS.get(ws.title, ws.title.replace("-", " / ")))
        notes += n

    for b in bugs:
        b["captured_on"] = captured
    for n in notes:
        n["captured_on"] = captured

    # ── the report ──
    sections = ["active", "resolved", "closed", "removed"]
    lines = []
    for ws in trackers:
        mine = [b for b in bugs if b["product"] == ws.title]
        counts = "  ".join(f"{s} {sum(1 for b in mine if b['section'] == s)}" for s in sections)
        label = LABELS.get(ws.title)
        lines.append(f"{ws.title:<24}{counts}" + ("" if label else "   (new tab — no label, shown as its tab name)"))
    totals = "  ".join(f"{s} {sum(1 for b in bugs if b['section'] == s)}" for s in sections)

    header = [
        "-- ═══════════════════════════════════════════════════════════════════════════",
        "--  TECHNICAL BUGS BACKLOG — generated by tools/build-tech-bugs.py. Do not edit;",
        "--  rebuild from the workbook and run the whole file.",
        "--",
        f"--  Source:  {book.name}, saved {captured.isoformat()}",
        "--",
    ] + [f"--    {l}" for l in lines] + [
        f"--    {'TOTAL':<24}{totals}",
        f"--    notes: {len(notes)}",
        "--",
        "--  Skipped: " + ", ".join(skipped),
        "--",
        "--  Safe to re-run. The rows are replaced inside one transaction, so a reader",
        "--  never sees the table empty between one week's list and the next.",
        "-- ═══════════════════════════════════════════════════════════════════════════",
        "",
    ]

    body = ["begin;", "", SCHEMA, "-- ── this week's list ──────────────────────────────────────────────────────",
            "delete from public.tech_bugs;", "delete from public.tech_bug_notes;", ""]
    body.append(f"insert into public.tech_bugs ({', '.join(BUG_COLS)}) values")
    body.append(",\n".join("  (" + ", ".join(sql(b[c]) for c in BUG_COLS) + ")" for b in bugs) + ";")
    if notes:
        body.append("")
        body.append(f"insert into public.tech_bug_notes ({', '.join(NOTE_COLS)}) values")
        body.append(",\n".join("  (" + ", ".join(sql(n[c]) for c in NOTE_COLS) + ")" for n in notes) + ";")
    body += ["",
             "-- ── this week, into the history ──────────────────────────────────────────",
             "-- Worked out from the rows just loaded, so the history can never disagree",
             "-- with the list. This week's rows first go, then come back; no other week",
             "-- is read or written.",
             f"delete from public.tech_bug_history where captured_on = {sql(captured)};",
             "insert into public.tech_bug_history",
             "       (captured_on, product, product_label, product_order, section,",
             "        bugs, scored, score_total, top_score)",
             "select captured_on, product, product_label, product_order, section,",
             "       count(*), count(score), coalesce(sum(score), 0), max(score)",
             "  from public.tech_bugs",
             " group by captured_on, product, product_label, product_order, section;",
             ]
    body += ["", "commit;", "",
             "-- ── check it ───────────────────────────────────────────────────────────────",
             "-- One row per tracker, with its counts per section. They should match the",
             "-- header of this file.",
             "select product_label,",
             "       count(*) filter (where section = 'active')   as active,",
             "       count(*) filter (where section = 'resolved') as resolved_this_week,",
             "       count(*) filter (where section = 'closed')   as closed,",
             "       count(*) filter (where section = 'removed')  as removed,",
             "       max(captured_on)                             as tracker_saved",
             "  from public.tech_bugs",
             " group by product_label, product_order",
             " order by product_order;", "",
             "-- The weeks on record, for the line graphs. One row per workbook run.",
             "select captured_on                                              as week,",
             "       sum(bugs) filter (where section = 'active')              as open_bugs,",
             "       sum(score_total) filter (where section = 'active')       as open_weighted_score,",
             "       sum(bugs) filter (where section = 'closed')              as closed,",
             "       sum(bugs) filter (where section = 'removed')             as removed,",
             "       sum(bugs) filter (where section = 'resolved')            as resolved_this_week",
             "  from public.tech_bug_history",
             " group by captured_on",
             " order by captured_on;", ""]

    OUT.write_text("\n".join(header + body), encoding="utf-8")
    print(f"wrote {OUT.relative_to(REPO)}  —  {book.name}, saved {captured.isoformat()}")
    for l in lines:
        print("  " + l)
    print(f"  {'TOTAL':<24}{totals}   notes {len(notes)}")
    print("  skipped: " + ", ".join(skipped))


if __name__ == "__main__":
    try:
        main()
    except BuildError as e:
        sys.exit(f"build-tech-bugs: {e}\nNothing was written.")
