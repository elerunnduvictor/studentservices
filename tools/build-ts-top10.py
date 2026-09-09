#!/usr/bin/env python3
"""Rebuild the Top 10 Tech Support list from a TS Product Tracker workbook.

    python tools/build-ts-top10.py "TS Product Tracker.xlsx"

── What it takes from the workbook ──

The tracker has one sheet per product, and every sheet's priority column
restarts at 1, so there is no ranked list of ten anywhere in it. The ten are
the priority-1 row from each product: one per product, covering the whole
estate rather than whichever product has the most bugs open this month. A
sheet with no priority column contributes its first issue row.

── What it does with what was there before ──

The list does not expire. It stands until a newer tracker replaces it, and
this is the part worth getting right when that happens:

    still on the new list      stays where it is
    no longer on it            moves to `retired`, dated
    back on it after a spell   leaves `retired` again

`retired` is what Backlog's tech support tab shows. Doing this by hand is how
an issue quietly disappears from both tabs, which is why it is a script.

Identity is product plus issue title, normalised. Bug numbers would be the
better key and cannot be: some rows have none, and others carry two.
"""
import datetime
import json
import re
import subprocess
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
TARGET = REPO / "emerging-issues" / "js" / "tech-support-top10.js"

# Sheets that are not a product's issue list.
NOT_PRODUCTS = {"product owners", "kbs", "transcripts"}

FIELDS = ("issue", "bug", "summary", "scope", "impact", "eta", "status")
MONTHS = ["January", "February", "March", "April", "May", "June",
          "July", "August", "September", "October", "November", "December"]
SHORT = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()


def clean(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return "%d %s %d" % (v.day, SHORT[v.month - 1], v.year)
    return re.sub(r"\s+", " ", str(v)).strip()


def key_of(rec):
    """What makes two rows the same issue, across two versions of the tracker."""
    return (rec.get("product", "").strip().lower(),
            re.sub(r"\s+", " ", rec.get("issue", "")).strip().lower())


def read_workbook(path):
    """The priority-1 row from every product sheet, in the workbook's order."""
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is needed to read the workbook:  pip install openpyxl")

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        if ws.title.strip().lower() in NOT_PRODUCTS:
            continue

        # The header is not always the first row — some sheets carry a banner
        # or a "last updated" line above it. Find the row that names a column
        # "Issue" and treat that as the header.
        rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 160),
                                 max_col=24, values_only=True))
        head = None
        for i, row in enumerate(rows):
            if any(clean(v).lower() == "issue" for v in row):
                head = i
                break
        if head is None:
            continue

        hdr = [clean(v).lower() for v in rows[head]]

        def col(*names):
            for n in names:
                for i, h in enumerate(hdr):
                    if h == n or h.startswith(n):
                        return i
            return None

        at = {
            "pri": col("priority", "p#"),
            "issue": col("issue"),
            "bug": col("bug#", "ics#", "bug /"),
            "summary": col("summary"),
            "scope": col("scope"),
            "impact": col("impact"),
            "eta": col("current eta", "eta"),
            "status": col("current status", "notes"),
        }

        best = None
        for row in rows[head + 1:]:
            get = lambda k: clean(row[at[k]]) if at[k] is not None and at[k] < len(row) else ""
            if not get("issue"):
                continue
            cand = {k: get(k) for k in at}
            if best is None:
                best = cand                       # no priority column: first row
            if re.fullmatch(r"1(\.0)?", cand["pri"]):
                best = cand
                break
        if not best:
            continue

        rec = {"product": ws.title.strip()}
        for f in FIELDS:
            v = best.get(f, "")
            if v and v.upper() not in ("N/A", "NA", "TBD"):
                rec[f] = v
        out.append(rec)
    return out


def read_existing():
    """The list currently on the page, as {issues, retired}."""
    if not TARGET.exists():
        return {"issues": [], "retired": []}
    script = (
        "global.window={};"
        "const fs=require('fs');"
        f"eval(fs.readFileSync({json.dumps(str(TARGET))},'utf8'));"
        "const d=window.TECH_SUPPORT_TOP10||{};"
        "process.stdout.write(JSON.stringify({issues:d.issues||[],retired:d.retired||[]}));"
    )
    r = subprocess.run(["node", "-e", script], capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0 or not r.stdout.strip():
        sys.exit("Could not read the current list:\n" + (r.stderr or ""))
    return json.loads(r.stdout)


def merge(previous, fresh, when):
    """The new ten, and what has fallen off the old ones.

    Order in `retired` is newest-dropped first, so the tab reads as a history
    rather than as whatever order the sheets happened to be in.
    """
    fresh_keys = {key_of(r) for r in fresh}
    prev_keys = {key_of(r) for r in previous["issues"]}

    # Back on the ten: it stops being retired.
    retired = [r for r in previous["retired"] if key_of(r) not in fresh_keys]

    # Off the ten: it becomes retired, dated, and keeps everything it had.
    dropped = []
    for r in previous["issues"]:
        if key_of(r) in fresh_keys:
            continue
        rec = dict(r)
        rec["dropped"] = when
        dropped.append(rec)

    added = [r for r in fresh if key_of(r) not in prev_keys]
    return fresh, dropped + retired, dropped, added


def render(issues, retired, today):
    label = "%d %s %d" % (today.day, MONTHS[today.month - 1], today.year)
    one = lambda r: "    " + json.dumps(r, ensure_ascii=False)
    body = ",\n".join(one(r) for r in issues)
    ret = ",\n".join(one(r) for r in retired)
    return f'''/* ════════ TOP 10 TECH SUPPORT ISSUES ════════

   The highest-priority open issue for each of the products Technical Support
   tracks, taken from the "TS Product Tracker" workbook.

   ── Why these ──

   The tracker has one sheet per product and each sheet has its own priority
   column that starts again at 1, so there is no single ranked list anywhere in
   it. These are the priority-1 rows: one per product, which covers the whole
   estate rather than whichever product happens to have the most bugs open at
   the moment. The order is the workbook's own sheet order.

   ── Refreshing it, and what happens to what falls off ──

   Send a fresh copy of the tracker and tools/build-ts-top10.py regenerates
   this file from it. The list does not expire on its own: it stands until a
   newer tracker replaces it.

   When one arrives, the two lists are compared. Anything still on the top ten
   stays where it is. Anything no longer on it moves to `retired`, which is
   what Backlog's tech support tab shows, with the date it dropped off. An
   issue that returns to the ten leaves `retired` again.

   Identity is product plus issue title — bug numbers are missing on some rows
   and doubled on others, so they cannot be the key.

   Generated {label}. Nothing here names a student — product, bug number,
   symptom, scope and ETA only — but it is internal operational detail, and it
   sits behind the gate in middleware.js once HUB_GATE is switched on.

   Do not hand-edit. Re-run the script.
   ════════════════════════════════════════════════════════ */

window.TECH_SUPPORT_TOP10 = {{
  captured: "{today.isoformat()}",
  // Spelled out as well as dated. "{today.isoformat()}" parsed as a Date is UTC
  // midnight, which prints as the day before anywhere west of Greenwich — the
  // page would have claimed the list was a day older than it is.
  capturedLabel: "{label}",
  source: "TS Product Tracker",

  /* Off the ten, still worth being able to find. Each keeps the date a later
     tracker stopped listing it. */
  retired: [
{ret}
  ],

  issues: [
{body}
  ],
}};
'''


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    book = Path(sys.argv[1])
    if not book.exists():
        sys.exit(f"No such workbook: {book}")

    today = datetime.date.today()
    when = "%d %s %d" % (today.day, MONTHS[today.month - 1], today.year)

    fresh = read_workbook(book)
    if not fresh:
        sys.exit("No product sheets with an Issue column — is this the right workbook?")

    previous = read_existing()
    issues, retired, dropped, added = merge(previous, fresh, when)

    TARGET.write_text(render(issues, retired, today), encoding="utf-8")

    print(f"{TARGET.relative_to(REPO)}")
    print(f"  on the list : {len(issues)}")
    print(f"  retired     : {len(retired)}")
    for r in dropped:
        print(f"    off  {r['product']}: {r['issue'][:56]}")
    for r in added:
        print(f"    new  {r['product']}: {r['issue'][:56]}")
    print("\nDelete the workbook when you are done with it.")


if __name__ == "__main__":
    main()
